#!/usr/bin/env python3
"""
Voucher standardisation and metadata curation (Step 5)
======================================================

1. ``standardize_vouchers``: one voucher per specimen. GenBank submitters write the same voucher
   differently per record (``COI_ZMMU_MSU_WS2585`` / ``28S_ZMMU_WS2585`` / ``WS2585``,
   ``ZMMU:WS30980`` / ``ZMMU_WS30980``). Each written form is normalised (gene names removed,
   separators unified), the most complete form is kept and written in the INSDC format
   ``[<institution-code>:[<collection-code>:]]<specimen_id>`` (``ZMMU:MSU:WS2585``, ``ZMMU:WS30980``).
   The organize step writes it to ``voucher_standardized`` next to ``voucher_as_submitted``
   (all forms as found in GenBank).

2. ``apply_updates`` / ``g2t-curate``: GenBank metadata is often outdated (coordinates, species names,
   publications added after submission). A user table with one row per correction updates the
   matrix. Rows are located by ``accession`` (any gene column; version optional) and/or ``voucher``
   (any written form; optionally narrowed with ``organism_match``). Empty cells change nothing.
   Every change is logged (old -> new, matched by); rows that match nothing, several rows, or
   contradict each other are reported and not applied. The input matrix is never modified.

3. ``apply_user_table``: the user's own table (any column names, Chinese or English) updates the
   matrix. Columns are recognised automatically (voucher/标本号, accession/登录号 (several gene
   columns allowed), species/拉丁名, 纬度 + 经度 -> lat_lon, 采集日期 -> DD-Mon-YYYY, title/作者/期刊 ...);
   ``--map`` overrides; other columns are added as ``user:<name>``.

CLI:
    g2t-curate -m matrix_A.xlsx -u my_table_B.xlsx -o curated_C.xlsx        # A + B -> C
    g2t-curate -m matrix_A.xlsx -u B.xlsx -o C.xlsx --map "编号=voucher"     # force a column
    g2t-curate -m matrix_A.xlsx --template corrections.xlsx                  # or a pre-filled template
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

import pandas as pd

GENE_TOKENS = (r"COI|CO1|COX1|COII|COX2|COIII|COX3|CYTB|COB|ND\d|NAD\d|ATP\d|12S|16S|18S|28S|5\.8S|"
               r"H3|ITS1?|ITS2|EF1A?|EF1ALPHA|RRNL|RRNS|LSU|SSU|MTDNA|MITO")
_PREFIX = re.compile(rf"^(?:{GENE_TOKENS})[\s_\-:]+", re.I)
_SUFFIX = re.compile(rf"[\s_\-:]+(?:{GENE_TOKENS})$", re.I)
KEY_COLUMNS = ["row_id", "voucher", "accession", "organism_match", "user_row"]
BASELINE_SHEET = "baseline (do not edit)"
TEMPLATE_FIELDS = ["organism", "lat_lon", "geo_loc_name", "country", "collection_date", "collected_by",
                   "identified_by", "Ref1Authors", "Ref1Title", "Ref1Journal", "curation_source"]
DEFAULT_GENES = ["mtgenome", "coi", "16s", "12s", "cob", "cox2", "cox3", "18s", "28s", "its1-its2",
                 "18-28s", "ef-1", "h3"]


# --------------------------------------------------------------------------- vouchers

def normalize_voucher(raw) -> str:
    """'COI_ZMMU_MSU_WS2585' -> 'ZMMU_MSU_WS2585'; 'ZMMU:WS30980' -> 'ZMMU_WS30980'."""
    v = "" if raw is None or (isinstance(raw, float) and pd.isna(raw)) else str(raw).strip()
    if not v or v.lower() == "nan":
        return ""
    for _ in range(2):
        v = _PREFIX.sub("", v)
        v = _SUFFIX.sub("", v)
    v = re.sub(r"[\s:]+", "_", v)
    v = re.sub(r"_+", "_", v).strip("_")
    return v


def _tokens(v: str) -> list[str]:
    return [t for t in re.split(r"[_\-]", v.upper()) if t]


def _contains(big: str, small: str) -> bool:
    """small's tokens appear in big, in order (e.g. WS2585 in ZMMU_MSU_WS2585)."""
    b, s = _tokens(big), _tokens(small)
    i = 0
    for t in b:
        if i < len(s) and t == s[i]:
            i += 1
    return i == len(s)


def to_insdc(norm: str) -> str:
    """Normalised voucher -> INSDC /specimen_voucher form "[<institution-code>:[<collection-code>:]]<specimen_id>".
    Leading letters-only tokens are the codes, the rest is the specimen id:
    ZMMU_WS30980 -> ZMMU:WS30980; ZMMU_MSU_WS12387_XZ5022 -> ZMMU:MSU:WS12387_XZ5022; WS0397 -> WS0397.
    With more than two code tokens the form is left unchanged (cannot tell institution from collection)."""
    if not norm:
        return ""
    toks = norm.split("_")
    codes = []
    for t in toks[:-1]:
        if re.fullmatch(r"[A-Za-z]{2,10}", t):
            codes.append(t)
        else:
            break
    if not codes or len(codes) > 2:
        return norm
    return ":".join(codes + ["_".join(toks[len(codes):])])


def standardize_vouchers(values) -> tuple[str, bool]:
    """-> (standardised voucher in INSDC form, consistent). The most complete normalised form is kept
    (see to_insdc); consistent is False when some form is not contained in it (more than a
    prefix/separator difference)."""
    forms: list[str] = []
    for raw in values:
        for part in str(raw if raw is not None else "").split(";"):
            n = normalize_voucher(part)
            if n and n not in forms:
                forms.append(n)
    if not forms:
        return "", True
    best = max(forms, key=lambda f: (len(_tokens(f)), len(f), -forms.index(f)))
    return to_insdc(best), all(_contains(best, f) for f in forms)


# --------------------------------------------------------------------------- curation

def _acc(a) -> str:
    return str(a).strip().split(".")[0].upper()


def _gene_cols(matrix: pd.DataFrame, gene_columns=None) -> list[str]:
    return [g for g in (gene_columns or DEFAULT_GENES) if g in matrix.columns]


def _voucher_forms(row: pd.Series) -> set[str]:
    forms = set()
    for c in ("voucher_standardized", "voucher_as_submitted", "specimen_voucher", "isolate"):
        if c in row and str(row[c]).strip() and str(row[c]).lower() != "nan":
            for part in str(row[c]).split(";"):
                n = normalize_voucher(part).upper()
                if n:
                    forms.add(n)
    return forms


def _blank(v) -> bool:
    return v is None or (isinstance(v, float) and pd.isna(v)) or str(v).strip() == "" or str(v).lower() == "nan"


def apply_updates(matrix: pd.DataFrame, updates: pd.DataFrame, gene_columns=None,
                  key_column: str = "species_voucher_new",
                  baseline: pd.DataFrame | None = None) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """-> (curated matrix, change log, problems). See module docstring.

    ``baseline``: the values a pre-filled template was made from (rows joined on ``row_id``). With it,
    only cells the user edited count as corrections, so unedited template cells never overwrite values
    that GenBank updated later; an edited cell whose GenBank value also changed is reported, not applied."""
    base_map: dict[str, dict] = {}
    if baseline is not None and "row_id" in baseline.columns:
        base_map = {str(r["row_id"]).strip(): r for _, r in baseline.fillna("").astype(str).iterrows()}
    out = matrix.copy().astype(object)
    genes = _gene_cols(out, gene_columns)
    acc_index: dict[str, set[int]] = {}
    for i, row in out.iterrows():
        for g in genes:
            for a in str(row[g]).split(";"):
                if a.strip() and a.strip().lower() != "nan":
                    acc_index.setdefault(_acc(a), set()).add(i)
    forms = {i: _voucher_forms(row) for i, row in out.iterrows()}
    if "curated_fields" not in out.columns:
        out["curated_fields"] = ""
    log, problems = [], []
    upd = updates.copy()
    upd.columns = [str(c).strip() for c in upd.columns]
    fields = [c for c in upd.columns if c not in KEY_COLUMNS]
    for c in fields:
        if c not in out.columns:
            out[c] = ""
            if c not in TEMPLATE_FIELDS and not c.startswith("user:"):
                problems.append({"update_row": "", "keys": "", "problem": f"new column '{c}' added to the matrix"})
    for n, r in upd.iterrows():
        acc = "" if _blank(r.get("accession")) else str(r.get("accession")).strip()
        vou = "" if _blank(r.get("voucher")) else str(r.get("voucher")).strip()
        org = "" if _blank(r.get("organism_match")) else str(r.get("organism_match")).strip()
        keys = "; ".join(x for x in (f"accession={acc}" if acc else "", f"voucher={vou}" if vou else "",
                                     f"organism_match={org}" if org else "") if x)
        rownum = n + 2                                   # spreadsheet row (header = 1)
        if not acc and not vou:
            problems.append({"update_row": rownum, "keys": keys, "problem": "no accession or voucher given"})
            continue
        hit_a = None
        if acc:
            per_acc = [acc_index.get(_acc(x), set()) for x in re.split(r"[;,\s]+", acc) if x.strip()]
            found = [h for h in per_acc if h]
            if len({frozenset(h) for h in found}) > 1:
                problems.append({"update_row": rownum, "keys": keys,
                                 "problem": "accessions point to different matrix rows; not applied"})
                continue
            hit_a = found[0] if found else set()
        q = normalize_voucher(vou).upper()
        hit_v = {i for i, f in forms.items() if any(_contains(x, q) or _contains(q, x) for x in f)} if vou else None
        if org:
            keep = {i for i in out.index if str(out.at[i, "organism"]).strip() == org} if "organism" in out else set()
            hit_a = hit_a & keep if hit_a is not None else None
            hit_v = hit_v & keep if hit_v is not None else None
        if hit_a is not None and hit_v is not None and hit_a != hit_v:
            problems.append({"update_row": rownum, "keys": keys,
                             "problem": "accession and voucher point to different rows; not applied"})
            continue
        hits = hit_a if hit_a is not None else hit_v
        if not hits:
            problems.append({"update_row": rownum, "keys": keys, "problem": "no matrix row matches; not applied"})
            continue
        if len(hits) > 1:
            names = ", ".join(str(out.at[i, key_column]) for i in sorted(hits)) if key_column in out else ""
            problems.append({"update_row": rownum, "keys": keys,
                             "problem": f"matches {len(hits)} matrix rows ({names}); add organism_match or "
                                        "use an accession; not applied"})
            continue
        i = next(iter(hits))
        by = "accession" if acc else "voucher"
        if "user_row" in upd.columns and not _blank(r.get("user_row")):
            rownum = r["user_row"]
        base = base_map.get(str(r.get("row_id", "")).strip()) if base_map else None
        for c in fields:
            new = r[c]
            if _blank(new):
                continue
            old = out.at[i, c]
            new = str(new).strip()
            cur = "" if _blank(old) else str(old).strip()
            if base is not None and c in base.index:
                was = str(base[c]).strip()
                if new == was:
                    continue                      # not edited by the user
                if cur != was:
                    problems.append({"update_row": rownum, "keys": keys,
                                     "problem": f"{c}: GenBank value changed since the template was made "
                                                f"('{was}' -> '{cur}'); your value '{new}' not applied"})
                    continue
            if cur == new:
                continue
            out.at[i, c] = new
            done = [x for x in str(out.at[i, "curated_fields"]).split("; ") if x]
            if c not in done:
                out.at[i, "curated_fields"] = "; ".join(done + [c])
            log.append({"specimen": out.at[i, key_column] if key_column in out else i, "field": c,
                        "old_value": "" if _blank(old) else old, "new_value": new, "matched_by": by,
                        "keys": keys, "update_row": rownum})
    log_df = pd.DataFrame(log, columns=["specimen", "field", "old_value", "new_value", "matched_by", "keys",
                                        "update_row"])
    prob_df = pd.DataFrame(problems, columns=["update_row", "keys", "problem"])
    return out, log_df, prob_df


# --------------------------------------------------------------------------- the user's own table

_ACC_VALUE = re.compile(r"^[A-Z]{1,2}_?\d{5,8}(\.\d+)?$")
# (target, name patterns); the first matching rule wins, so specific names come first
_ALIASES = [
    ("lat_lon", [r"^lat_?lon$", r"coordinates?", r"经纬度"]),
    ("latitude", [r"^lat$", r"latitude", r"纬度"]),
    ("longitude", [r"^lon$", r"^long$", r"^lng$", r"longitude", r"经度"]),
    ("accession", [r"accession", r"genbank", r"登录号", r"^acc(\.|_no)?$"]),
    ("voucher", [r"voucher", r"catalog", r"catalogue", r"museum.?(no|number|id)", r"specimen.?(no|number|id)?$",
                 r"凭证", r"标本号", r"标本编号", r"馆藏号"]),
    ("organism", [r"^species$", r"scientific.?name", r"^organism$", r"^taxon$", r"拉丁名", r"^学名$", r"^物种$",
                  r"^种名$"]),
    ("geo_loc_name", [r"geo_loc_name", r"locality", r"^location$", r"^site$", r"采集地", r"产地", r"^地点$"]),
    ("country", [r"^country$", r"国家"]),
    ("collection_date", [r"collection.?date", r"^date$", r"采集日期", r"采集时间", r"^日期$"]),
    ("Ref1Authors", [r"^authors?$", r"作者"]),
    ("Ref1Journal", [r"^journal$", r"期刊"]),
    ("Ref1Title", [r"^title$", r"publication", r"reference", r"^paper$", r"^article$", r"题目", r"文献", r"论文"]),
    ("collected_by", [r"collector", r"collected.?by", r"采集人"]),
    ("identified_by", [r"identifier", r"identified.?by", r"determiner", r"鉴定人"]),
]
_SINGLE = {"voucher", "organism", "latitude", "longitude", "lat_lon", "geo_loc_name", "country",
           "collection_date", "Ref1Title", "Ref1Authors", "Ref1Journal", "collected_by", "identified_by"}
_PREFER = re.compile(r"实测|measured|actual|final|corrected|校正|修正", re.I)
_MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def detect_columns(df: pd.DataFrame, column_map: dict | None = None) -> dict[str, str]:
    """User column -> target field ('voucher', 'accession', 'organism', 'latitude', ... or 'user:<name>').
    ``column_map`` ({user column: target}) overrides the automatic choice."""
    column_map = column_map or {}
    found: dict[str, str] = {}
    for col in df.columns:
        if col in column_map:
            found[col] = column_map[col]
            continue
        name = str(col).strip().lower()
        target = ""
        if "中文" not in name and "chinese" not in name:
            for t, pats in _ALIASES:
                if any(re.search(p, name) for p in pats):
                    target = t
                    break
        if not target:
            vals = [str(v).strip() for v in df[col] if str(v).strip() and str(v).lower() != "nan"]
            if vals and sum(bool(_ACC_VALUE.match(v)) for v in vals) / len(vals) >= 0.6:
                target = "accession"
        found[col] = target or f"user:{col}"
    # one column per single-valued field: prefer 'measured/实测/corrected' names, the rest stay user columns
    for t in _SINGLE:
        cols = [c for c, v in found.items() if v == t and c not in column_map]
        if len(cols) > 1:
            keep = next((c for c in cols if _PREFER.search(str(c))), cols[0])
            for c in cols:
                if c != keep:
                    found[c] = f"user:{c}"
    return found


def _num(v) -> float | None:
    try:
        return float(str(v).strip())
    except ValueError:
        return None


def genbank_lat_lon(lat, lon) -> str:
    """Decimal degrees -> GenBank lat_lon '66.55 N 33.10 E' (the digits as given)."""
    la, lo = _num(lat), _num(lon)
    if la is None or lo is None:
        return ""
    def part(raw, v, pos, neg):
        txt = str(raw).strip().lstrip("+-")
        return f"{txt} {pos if v >= 0 else neg}"
    return f"{part(lat, la, 'N', 'S')} {part(lon, lo, 'E', 'W')}"


def genbank_date(v) -> str:
    """'20190612' / '2019-06-12' / '2019/6/12' / Excel datetime -> '12-Jun-2019'; '2019-06' -> 'Jun-2019'."""
    s = str(v).strip()
    if not s or s.lower() == "nan":
        return ""
    m = re.fullmatch(r"(\d{4})[-/.]?(\d{1,2})[-/.]?(\d{1,2})(?:[ T]00:00:00(?:\.0+)?)?", s)
    if m and len(s) >= 8:
        y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if 1 <= mo <= 12 and 1 <= d <= 31:
            return f"{d:02d}-{_MONTHS[mo - 1]}-{y}"
    m = re.fullmatch(r"(\d{4})[-/.](\d{1,2})", s)
    if m and 1 <= int(m.group(2)) <= 12:
        return f"{_MONTHS[int(m.group(2)) - 1]}-{m.group(1)}"
    return s


def apply_user_table(matrix: pd.DataFrame, user: pd.DataFrame, key_column: str = "species_voucher_new",
                     gene_columns=None, column_map: dict | None = None
                     ) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Update the NCBI matrix from the user's own table (any column names).
    -> (updated matrix, change log, problems, column mapping)."""
    user = user.fillna("").astype(str)
    mapping = detect_columns(user, column_map)
    acc_cols = [c for c, t in mapping.items() if t == "accession"]
    inv = {t: c for c, t in mapping.items() if t in _SINGLE}
    rows = []
    for n, r in user.iterrows():
        row = {"user_row": n + 2,
               "accession": "; ".join(str(r[c]).strip() for c in acc_cols if str(r[c]).strip()),
               "voucher": str(r[inv["voucher"]]).strip() if "voucher" in inv else ""}
        if "latitude" in inv and "longitude" in inv:
            row["lat_lon"] = genbank_lat_lon(r[inv["latitude"]], r[inv["longitude"]])
        elif "lat_lon" in inv:
            row["lat_lon"] = str(r[inv["lat_lon"]]).strip()
        for t in ("organism", "geo_loc_name", "country", "Ref1Title", "Ref1Authors", "Ref1Journal",
                  "collected_by", "identified_by"):
            if t in inv:
                row[t] = str(r[inv[t]]).strip()
        if "collection_date" in inv:
            row["collection_date"] = genbank_date(r[inv["collection_date"]])
        for c, t in mapping.items():
            if t.startswith("user:"):
                row[t] = str(r[c]).strip()
        rows.append(row)
    upd = pd.DataFrame(rows)
    out, log, problems = apply_updates(matrix, upd, gene_columns=gene_columns, key_column=key_column)
    map_df = pd.DataFrame([{"user_column": c, "used_as": t if not t.startswith("user:") else "(added as a new column)",
                            "note": "combined into lat_lon (GenBank format)" if t in ("latitude", "longitude")
                            else ("converted to GenBank date format" if t == "collection_date" else
                                  ("located the specimen" if t in ("voucher", "accession") else ""))}
                           for c, t in mapping.items()])
    return out, log, problems, map_df


def make_template(matrix: pd.DataFrame, gene_columns=None, fields=None) -> pd.DataFrame:
    """One row per specimen, pre-filled with the current values, for the user to correct."""
    genes = _gene_cols(matrix, gene_columns)
    rows = []
    for _, r in matrix.iterrows():
        acc = next((str(r[g]).split(";")[0].strip() for g in genes
                    if str(r[g]).strip() and str(r[g]).lower() != "nan"), "")
        voucher = r.get("voucher_standardized", "")
        row = {"row_id": len(rows) + 1, "voucher": "" if _blank(voucher) else voucher, "accession": acc,
               "organism_match": r.get("organism", "")}
        for f in fields or TEMPLATE_FIELDS:
            v = r.get(f, "")
            row[f] = "" if _blank(v) else v
        rows.append(row)
    return pd.DataFrame(rows)


def _read(path: str, sheet=0) -> pd.DataFrame:
    p = Path(path)
    if p.suffix.lower() in (".xlsx", ".xls"):
        return pd.read_excel(p, dtype=str, sheet_name=sheet).fillna("")
    sep = "\t" if p.suffix.lower() in (".tsv", ".txt") else ","
    return pd.read_csv(p, dtype=str, sep=sep, keep_default_na=False)


def write_template(template: pd.DataFrame, path: str) -> None:
    """xlsx: sheets 'corrections' + baseline; csv: <name>.csv + <name>.baseline.csv."""
    p = Path(path)
    if p.suffix.lower() == ".xlsx":
        with pd.ExcelWriter(p) as w:
            template.to_excel(w, sheet_name="corrections", index=False)
            template.to_excel(w, sheet_name=BASELINE_SHEET, index=False)
    else:
        template.to_csv(p, index=False)
        template.to_csv(p.with_name(p.stem + ".baseline.csv"), index=False)


def read_corrections(path: str) -> tuple[pd.DataFrame, pd.DataFrame | None]:
    """-> (corrections, baseline or None)."""
    p = Path(path)
    if p.suffix.lower() in (".xlsx", ".xls"):
        sheets = pd.read_excel(p, dtype=str, sheet_name=None)
        names = list(sheets)
        upd = sheets["corrections"] if "corrections" in sheets else sheets[names[0]]
        base = sheets.get(BASELINE_SHEET)
        return upd.fillna(""), (base.fillna("") if base is not None else None)
    b = p.with_name(p.stem + ".baseline.csv")
    return _read(path), (_read(str(b)) if b.exists() else None)


_ACC_CELL = re.compile(r"^[A-Z]{1,6}_?\d{5,}(\.\d+)?$")


def matrix_gene_columns(matrix: pd.DataFrame) -> list[str]:
    """Gene columns = columns whose filled cells are all accession(.version)s."""
    out = []
    for c in matrix.columns:
        vals = [str(v).strip() for v in matrix[c] if not _blank(v)]
        if vals and all(_ACC_CELL.match(v) for v in vals):
            out.append(c)
    return out


def read_matrix(path: str) -> pd.DataFrame:
    """The NCBI matrix: a csv/tsv, or an Excel whose sheet 'Matrix' (else the first sheet) holds it."""
    p = Path(path)
    if p.suffix.lower() in (".xlsx", ".xls"):
        names = pd.ExcelFile(p).sheet_names
        sheet = next((n for n in names if n.strip().lower() in ("matrix", "matrix (curated)")), names[0])
        return _read(path, sheet)
    return _read(path)


def matrix_key_column(matrix: pd.DataFrame) -> str:
    for c in ("specimen_key", "species_voucher_new", "species_voucher"):
        if c in matrix.columns:
            return c
    return matrix.columns[0]


def is_template(path: str) -> bool:
    """True for a g2t corrections template (row_id column or a baseline copy), False for the user's own table."""
    p = Path(path)
    if p.suffix.lower() in (".xlsx", ".xls"):
        names = pd.ExcelFile(p).sheet_names
        if BASELINE_SHEET in names:
            return True
        cols = pd.read_excel(p, nrows=0, sheet_name=names[0]).columns
    else:
        cols = _read(path).columns
    return "row_id" in cols


def write_curated_workbook(path: str, out: pd.DataFrame, log: pd.DataFrame, problems: pd.DataFrame,
                           original: pd.DataFrame, mapping: pd.DataFrame | None = None) -> None:
    """Excel C: Matrix (updated cells highlighted) / Changes / Problems / [Column mapping] / Matrix (GenBank)."""
    from openpyxl.styles import PatternFill
    from openpyxl.utils import get_column_letter
    sheets = {"Matrix": out, "Changes": log, "Problems": problems}
    if mapping is not None:
        sheets["Column mapping"] = mapping
    sheets["Matrix (GenBank)"] = original
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        for name, df in sheets.items():
            (df if len(df.columns) else pd.DataFrame({"note": ["none"]})).to_excel(w, sheet_name=name, index=False)
            ws = w.sheets[name]
            ws.freeze_panes = "B2"
            for j, c in enumerate(df.columns, 1):
                width = max([len(str(c))] + [len(str(v)) for v in df[c].head(200)])
                ws.column_dimensions[get_column_letter(j)].width = min(max(width + 2, 8), 50)
        if "curated_fields" in out.columns:
            fill = PatternFill("solid", fgColor="FFF2CC")
            ws = w.sheets["Matrix"]
            col = {c: j for j, c in enumerate(out.columns, 1)}
            for n, fields in enumerate(out["curated_fields"], 2):
                for f in [x.strip() for x in str(fields).split(";") if x.strip() and str(fields) != "nan"]:
                    if f in col:
                        ws.cell(row=n, column=col[f]).fill = fill


def parse_map(items) -> dict:
    """['标本号=voucher', 'Depth=user:Depth'] -> {column: target}."""
    out = {}
    for it in items or []:
        if "=" not in it:
            raise ValueError(f"--map needs 'column=field', got {it!r}")
        k, v = it.split("=", 1)
        out[k.strip()] = v.strip()
    return out


def main(argv: list[str] | None = None) -> None:
    p = argparse.ArgumentParser(
        description="Update the NCBI matrix (A) from your own table (B) -> a corrected matrix (C)")
    p.add_argument("-m", "--matrix", required=True, help="matrix A: .xlsx (sheet 'Matrix') / .csv / .tsv")
    p.add_argument("-u", "--updates", help="table B: your own Excel/csv with any column names "
                                           "(or a g2t corrections template)")
    p.add_argument("-o", "--output", help="output C: .xlsx (Matrix / Changes / Problems / Column mapping / "
                                          "Matrix (GenBank)) or .csv (log and problems written next to it)")
    p.add_argument("--sheet", default=0, help="sheet of table B (name or 0-based number; default first)")
    p.add_argument("--map", action="append", metavar="COLUMN=FIELD",
                   help="force a column of B to a field, e.g. --map '编号=voucher' --map 'Lat=latitude'; "
                        "FIELD = voucher, accession, organism, latitude, longitude, lat_lon, geo_loc_name, "
                        "country, collection_date, Ref1Title, Ref1Authors, Ref1Journal, collected_by, "
                        "identified_by, or user:<new column>")
    p.add_argument("--key", help="matrix key column (default: specimen_key / species_voucher_new)")
    p.add_argument("--template", help="write a pre-filled corrections template (.csv or .xlsx) and exit")
    a = p.parse_args(argv)
    matrix = read_matrix(a.matrix)
    genes = matrix_gene_columns(matrix)
    key = a.key or matrix_key_column(matrix)
    if a.template:
        t = make_template(matrix, gene_columns=genes)
        write_template(t, a.template)
        print(f"Template: {a.template} ({len(t)} specimens). Edit the cells to correct, leave blank to keep.")
        return
    if not (a.updates and a.output):
        p.error("-u/--updates and -o/--output are required (or use --template)")
    mapping = None
    if is_template(a.updates):
        upd, base = read_corrections(a.updates)
        out, log, problems = apply_updates(matrix, upd, gene_columns=genes, key_column=key, baseline=base)
    else:
        sheet = int(a.sheet) if str(a.sheet).isdigit() else a.sheet
        table = _read(a.updates, sheet)
        out, log, problems, mapping = apply_user_table(matrix, table, key_column=key, gene_columns=genes,
                                                       column_map=parse_map(a.map))
    o = Path(a.output)
    if o.suffix.lower() == ".xlsx":
        write_curated_workbook(str(o), out, log, problems, matrix, mapping)
    else:
        out.to_csv(o, index=False)
        log.to_csv(o.with_name(o.stem + "_curation_log.csv"), index=False)
        problems.to_csv(o.with_name(o.stem + "_curation_problems.csv"), index=False)
        if mapping is not None:
            mapping.to_csv(o.with_name(o.stem + "_column_mapping.csv"), index=False)
    print(f"Curated: {o} ({len(log)} changes in {log['specimen'].nunique() if len(log) else 0} specimens; "
          f"{len(problems)} problems)")
    if mapping is not None:
        for _, r in mapping.iterrows():
            print(f"  {r['user_column']!s:<24} -> {r['used_as']}")


if __name__ == "__main__":
    main()
